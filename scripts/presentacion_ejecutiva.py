"""Prototipo de presentación ejecutiva en PowerPoint (fase 5 del roadmap: evaluación).

Arma un .pptx con lo que la corrida diaria ya dejó: la última fila de
Fact_Mercado_Macro (SELECT, solo lectura), su párrafo y sus comentarios de IA,
y los gráficos de Previews/. No scrapea, no llama a Gemini y no escribe nada
salvo el archivo de salida, que va a una carpeta temporal y nunca a Previews/.

Es un prototipo para decidir si vale la pena: no está conectado al pipeline.
La evaluación completa está en docs/evaluacion-powerpoint-y-streamlit.md.

Uso:
    python scripts/presentacion_ejecutiva.py                 # .pptx en una carpeta temporal
    python scripts/presentacion_ejecutiva.py --salida DIR
    python scripts/presentacion_ejecutiva.py --graficos DIR  # otra carpeta de gráficos
"""

import argparse
import sys
import tempfile
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import charts
import data_access
import email_report
import ia_generator
import transformations

# Los colores del mail
AZUL = RGBColor(0x2C, 0x3E, 0x50)
NARANJA = RGBColor(0xF3, 0x9C, 0x12)
GRIS = RGBColor(0x7F, 0x8C, 0x8D)
VERDE = RGBColor(0x1E, 0x84, 0x49)
ROJO = RGBColor(0xC0, 0x39, 0x2B)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

ANCHO, ALTO = Inches(13.333), Inches(7.5)  # 16:9


def _texto(slide, x, y, ancho, alto, texto, tamanio=14, color=AZUL, negrita=False, alinear=PP_ALIGN.LEFT):
    caja = slide.shapes.add_textbox(x, y, ancho, alto)
    marco = caja.text_frame
    marco.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.alignment = alinear
        run = parrafo.add_run()
        run.text = linea
        run.font.size = Pt(tamanio)
        run.font.bold = negrita
        run.font.color.rgb = color
    return caja


def _encabezado(prs, titulo: str, fecha: date):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    franja = slide.shapes.add_shape(1, 0, 0, ANCHO, Inches(0.9))  # 1 = rectángulo
    franja.fill.solid()
    franja.fill.fore_color.rgb = AZUL
    franja.line.fill.background()
    _texto(slide, Inches(0.4), Inches(0.15), Inches(10), Inches(0.6), titulo, 26, BLANCO, True)
    _texto(slide, Inches(10.4), Inches(0.25), Inches(2.6), Inches(0.5), f"{fecha:%d/%m/%Y}", 16, BLANCO,
           alinear=PP_ALIGN.RIGHT)
    return slide


def _tarjeta(slide, x, y, titulo: str, valor: str, detalle: str, color_detalle=GRIS):
    fondo = slide.shapes.add_shape(1, x, y, Inches(3.9), Inches(1.6))
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = RGBColor(0xEE, 0xF1, 0xF4)
    fondo.line.color.rgb = RGBColor(0xDF, 0xE4, 0xEA)
    _texto(slide, x + Inches(0.2), y + Inches(0.1), Inches(3.5), Inches(0.4), titulo.upper(), 11, GRIS, True)
    _texto(slide, x + Inches(0.2), y + Inches(0.45), Inches(3.5), Inches(0.6), valor, 28, AZUL, True)
    _texto(slide, x + Inches(0.2), y + Inches(1.1), Inches(3.5), Inches(0.4), detalle, 12, color_detalle)


def _variacion(hoy: float, ayer: float) -> tuple[str, RGBColor]:
    var = (hoy / ayer - 1) * 100
    return f"{var:+.2f}% contra la rueda anterior", (VERDE if var > 0 else ROJO if var < 0 else GRIS)


def armar(df, imagenes: dict[str, Path], salida: Path) -> Path:
    """`df` es Fact_Mercado_Macro más nuevo primero, como lo devuelve data_access.leer_historico."""
    hoy, ayer = df.iloc[0], df.iloc[1]
    fecha = date.fromisoformat(hoy["Fecha"])
    fwd_oficial, fwd_blue = transformations.forwards_fisher(df)
    resumen = email_report.resumen_ejecutivo(transformations.agregar_brechas_y_variaciones(df))
    comentarios = ia_generator.comentarios_por_grafico(hoy.get("ai_secciones"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO

    # 1. Portada
    portada = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = portada.shapes.add_shape(1, 0, 0, ANCHO, ALTO)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = AZUL
    fondo.line.fill.background()
    _texto(portada, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2), "Reporte Macroeconómico", 48, BLANCO, True)
    _texto(portada, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.8),
           "Tipos de cambio, riesgo país, tasas e inflación", 22, NARANJA)
    _texto(portada, Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.6), f"{fecha:%d/%m/%Y}", 20, BLANCO)

    # 2. Tablero
    tablero = _encabezado(prs, "Tablero del día", fecha)
    tarjetas = [
        ("Dólar blue", f"$ {hoy['TCV_Blue']:,.2f}", *_variacion(hoy["TCV_Blue"], ayer["TCV_Blue"])),
        ("Dólar MEP", f"$ {hoy['TCV_MEP']:,.2f}", *_variacion(hoy["TCV_MEP"], ayer["TCV_MEP"])),
        ("Billete BNA", f"$ {hoy['TCV_Billete']:,.2f}", *_variacion(hoy["TCV_Billete"], ayer["TCV_Billete"])),
        ("Riesgo país", f"{hoy['riesgo_pais']:,.0f} pts",
         f"{hoy['riesgo_pais'] - ayer['riesgo_pais']:+,.0f} pts contra la rueda anterior",
         ROJO if hoy["riesgo_pais"] > ayer["riesgo_pais"] else VERDE),
        ("BADLAR (TEA)", f"{hoy['bcra_tea']:.2f}%", f"FED (EFFR): {hoy['fed_tea']:.2f}%", GRIS),
        ("Forward oficial 3 meses", f"$ {fwd_oficial:,.2f}", f"Forward blue: $ {fwd_blue:,.2f}", GRIS),
    ]
    for i, (titulo, valor, detalle, color) in enumerate(tarjetas):
        x = Inches(0.45 + (i % 3) * 4.2)
        y = Inches(1.4 + (i // 3) * 2.0)
        _tarjeta(tablero, x, y, titulo, valor, detalle, color)
    diferencia = abs(resumen["ahorro_valor"])
    sentido = "más barato" if resumen["ahorro_valor"] > 0 else "más caro"
    _texto(tablero, Inches(0.45), Inches(5.6), Inches(12.4), Inches(1.2),
           f"Comprar USD {resumen['cantidad_usd']} al oficial con impuestos (solidario) sale $ {diferencia:,.2f} "
           f"{sentido} que en el blue: el solidario está {abs(resumen['brecha_solidario']):.2f}% "
           f"{'por debajo' if resumen['brecha_solidario'] > 0 else 'por encima'} del blue.",
           14, GRIS)

    # 3. Análisis de IA
    analisis = _encabezado(prs, "Análisis", fecha)
    parrafo = hoy["ai_paragraph"] if isinstance(hoy["ai_paragraph"], str) and hoy["ai_paragraph"] else "Sin análisis para este día."
    _texto(analisis, Inches(0.6), Inches(1.3), Inches(12.1), Inches(2.4), parrafo, 18, AZUL)
    _texto(analisis, Inches(0.6), Inches(6.8), Inches(12.1), Inches(0.4),
           "Generado con Gemini a partir de los datos del día. No es asesoramiento financiero.", 11, GRIS)

    # 4 a 6. Gráficos, con sus comentarios si los hay
    def diapositiva_grafico(titulo: str, nombres: list[str], cid: str | None):
        slide = _encabezado(prs, titulo, fecha)
        presentes = [imagenes[n] for n in nombres if n in imagenes]
        textos = comentarios.get(cid, []) if cid else []
        cuerpo = "\n\n".join(f"{t.upper()}\n{x}" for t, x in textos)

        if len(presentes) > 1:
            # Dos gráficos lado a lado
            for j, ruta in enumerate(presentes):
                slide.shapes.add_picture(str(ruta), Inches(0.45) + j * Inches(6.3), Inches(1.6), width=Inches(6.1))
            return
        if not presentes:
            _texto(slide, Inches(0.6), Inches(3), Inches(12), Inches(1), "Gráfico no disponible para este día.", 18, GRIS)
            return

        foto = slide.shapes.add_picture(str(presentes[0]), Inches(0.45), Inches(1.1), height=Inches(6.2))
        if foto.width > foto.height:
            # Apaisado (BTC): más bajo y centrado, con el comentario abajo
            foto.height, foto.width = int(foto.height * 0.75), int(foto.width * 0.75)
            foto.left = int((ANCHO - foto.width) / 2)
            if textos:
                _texto(slide, Inches(0.6), foto.top + foto.height + Inches(0.15), Inches(12.1), Inches(1.5), cuerpo, 13, AZUL)
        elif textos:
            # Vertical (tipos de cambio): el comentario a la derecha
            izquierda = foto.left + foto.width + Inches(0.4)
            _texto(slide, izquierda, Inches(1.2), ANCHO - izquierda - Inches(0.4), Inches(6), cuerpo, 13, AZUL)
        else:
            foto.left = int((ANCHO - foto.width) / 2)

    diapositiva_grafico("Tipos de cambio y riesgo país", [charts.TIPOS_DE_CAMBIO], "image1")
    diapositiva_grafico("Inflación y variaciones acumuladas", [charts.INFLACION, charts.VARIACIONES], None)
    diapositiva_grafico("Bitcoin", [charts.BTC], "image4")

    salida.mkdir(parents=True, exist_ok=True)
    ruta = salida / f"Reporte Macroeconomico {fecha:%Y-%m-%d}.pptx"
    prs.save(ruta)
    return ruta


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--salida", type=Path, help="Carpeta del .pptx (por defecto, una temporal)")
    parser.add_argument("--graficos", type=Path, default=RAIZ / "Previews", help="Carpeta de los .jpg")
    args = parser.parse_args()

    engine = data_access.crear_engine()
    df, _ = data_access.leer_historico(engine, respaldo_csv=False)
    engine.dispose()

    imagenes = {n: args.graficos / n for n in charts.ORDEN_EN_MAIL if (args.graficos / n).exists()}
    ruta = armar(df, imagenes, args.salida or Path(tempfile.mkdtemp(prefix="presentacion_")))
    print(f"Presentación: {ruta}")


if __name__ == "__main__":
    main()
