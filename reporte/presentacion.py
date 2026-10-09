"""Presentación ejecutiva en PowerPoint: nueve diapositivas con lo que deja la corrida del día.

1. Portada blanca con el logo de GlobalAIze, los links del autor y las tecnologías del proyecto.
2. Tablero: blue, MEP, billete, riesgo país, BADLAR y el forward de Fisher.
3. Tipos de cambio: el análisis de IA del día y los paneles de paralelas y oficiales.
4. Riesgo país.
5. Inflación y variaciones acumuladas.
6 y 7. Agregados monetarios y deuda, con el resumen del día y lo ideal.
8. Bitcoin, en fondo negro.
9. Fuentes y contacto, como el pie del mail.

Cada gráfico lleva su frase: la de Gemini cuando la hay y, si no, un resumen
calculado acá con los datos del día. Los colores siguen el criterio del mail,
mirado desde la macro: que el dólar o el riesgo país bajen es una buena noticia
(verde) y que suban, una mala (rojo).

La arma todos los días la etapa `presentacion` de pipeline.py, con los datos y
los textos de esa misma corrida, y queda en Previews/ con un nombre fijo
(ARCHIVO), que git ignora: la etapa previews la publica sola en la rama
reporte-ejecutivo, que se reemplaza entera cada día, así GitHub tiene siempre la
del último reporte sin acumular versiones. Para rearmarla a mano desde Supabase
está scripts/presentacion_ejecutiva.py.
"""

import os
import tempfile
from datetime import UTC, date, datetime
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from reporte import charts, email_report, ia_generator, indicadores, transformations
from reporte.scrapers import ambito, bcra, bna, dolarhoy, fed, finanzas, riesgo_pais

ARCHIVO = "Reporte Ejecutivo.pptx"

# Metadatos del archivo. La plantilla de python-pptx trae los de su autor
# (last_modified_by, un comentario y fechas de 2013): se reemplazan todos.
AUTOR = "Seguimiento Macroeconómico"

# Los colores del mail, los de GlobalAIze
# Los azules del logo de GlobalAIze (#081E40) y uno intermedio para los degradés
OSCURO = RGBColor(0x08, 0x1E, 0x40)
AZUL = RGBColor(0x16, 0x3A, 0x6B)
# Acentos con los colores de GlobalAIze, sin naranja: el gris azulado del logo para
# líneas y bordes, y un celeste claro del mismo tono para los rótulos sobre fondo oscuro
ACENTO = RGBColor(0x69, 0x7E, 0x91)
ACENTO_CLARO = RGBColor(0xA8, 0xC3, 0xDB)
GRIS = RGBColor(0x5D, 0x6D, 0x7E)
CLARO = RGBColor(0xEC, 0xF0, 0xF1)
FONDO_TARJETA = RGBColor(0xE7, 0xF2, 0xF8)  # el celeste claro del logo
VERDE = RGBColor(0x1E, 0x84, 0x49)
ROJO = RGBColor(0xC0, 0x39, 0x2B)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x00, 0x00, 0x00)
CELESTE = RGBColor(0x29, 0x80, 0xB9)
FONDO_SLIDE = RGBColor(0xF2, 0xF6, 0xFA)  # gris azulado muy claro, para que resalten las tarjetas blancas
BORDE_TARJETA = RGBColor(0xDC, 0xE6, 0xEF)
TARJETA_OSCURA = RGBColor(0x0B, 0x1B, 0x33)  # las tarjetas de la diapositiva negra de BTC
GRIS_LOGO = RGBColor(0x69, 0x7E, 0x91)  # el gris azulado del logo de GlobalAIze

ANCHO, ALTO = Inches(13.333), Inches(7.5)  # 16:9

SIN_ANALISIS = "Sin análisis para este día."

RECURSOS = Path(__file__).resolve().parent / "recursos"
ALIAS = "mauroemartinezmp"
MAIL = "mauro@globalaize.com"


def _icono_globalaize() -> Path:
    """El logo de GlobalAIze si está en recursos/; si no, el ícono provisorio."""
    logo = RECURSOS / "globalaize_logo.png"
    return logo if logo.exists() else RECURSOS / "globalaize.png"


# (ícono, rótulo, link): los mismos del cierre del mail
LINKS = [
    ("linkedin.png", "LinkedIn", "https://linkedin.com/in/mauroemartinez/"),
    ("github.png", "GitHub", "https://github.com/mauroemartinez/"),
    (None, "GlobalAIze", "https://www.globalaize.com"),
    ("linkedin.png", "GlobalAIze en LinkedIn", "https://www.linkedin.com/company/globalaize"),
]

# Las tecnologías del pie de la portada: (archivo en recursos/tecnologias, rótulo)
TECNOLOGIAS = [
    ("python", "Python"), ("playwright", "Playwright"), ("pandas", "pandas"), ("pydantic", "Pydantic"),
    ("sqlalchemy", "SQLAlchemy"), ("postgresql", "PostgreSQL"), ("supabase", "Supabase"), ("gemini", "Gemini"),
    ("matplotlib", "Matplotlib"), ("seaborn", "seaborn"), ("jinja", "Jinja"), ("powerpoint", "PowerPoint"),
    ("githubactions", "GitHub Actions"), ("pytest", "pytest"), ("fastapi", "FastAPI"), ("docker", "Docker"),
]

FUENTES = [
    ("Cotizaciones oficiales", bna.WEB_BNA),
    ("Cotizaciones paralelas", dolarhoy.WEB_DOLARHOY),
    ("Dólar MEP", ambito.WEB_MEP),
    ("Euro blue", ambito.WEB_EURO),
    ("Riesgo país y feriados", riesgo_pais.API_BASE),
    ("Tasa de la FED (EFFR)", fed.API_URL),
    ("BCRA: tasas, inflación, agregados, deuda y dólar mayorista", bcra.API_BASE),
    ("Deuda bruta del Tesoro (Secretaría de Finanzas)", finanzas.PAGINA),
    ("Bitcoin (Yahoo Finance)", "https://finance.yahoo.com/quote/BTC-USD"),
]


# ── Utilidades ───────────────────────────────────────────────────────────────

def _fecha(valor) -> date:
    """La Fecha de una fila: texto 'YYYY-MM-DD' en el histórico, date en la fila recién armada."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return date.fromisoformat(str(valor)[:10])


def _num(valor: float, decimales: int = 2) -> str:
    return indicadores.numero(valor, decimales)


def _signo(valor: float, decimales: int = 2) -> str:
    return f"{'+' if valor >= 0 else '-'}{_num(abs(valor), decimales)}"


def _texto(slide, x, y, ancho, alto, texto, tamanio=14, color=AZUL, negrita=False, alinear=PP_ALIGN.LEFT,
           link: str | None = None):
    caja = slide.shapes.add_textbox(x, y, ancho, alto)
    marco = caja.text_frame
    marco.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.alignment = alinear
        run = parrafo.add_run()
        run.text = linea
        run.font.size = Pt(tamanio)
        run.font.bold = negrita
        run.font.color.rgb = color
    if link:
        # El link va en la caja y no en el texto: en el texto, PowerPoint lo pinta con
        # el azul del tema, que sobre fondo oscuro no se lee
        caja.click_action.hyperlink.address = link
    return caja


def _tamanio_que_entra(bloques, ancho, alto, maximo: int = 16, minimo: int = 11) -> int:
    """El tamaño de letra más grande con el que las secciones entran en la caja.

    Es una estimación: el ancho promedio de una letra de Calibri es cerca de media
    vez su tamaño, y cada renglón ocupa 1,2 veces el tamaño más el espacio entre
    párrafos. Se queda un poco corta a propósito, para que nada se salga de la caja.
    """
    util_ancho = ancho / 12700 - 14.4  # EMU a puntos, menos los márgenes internos de la caja
    util_alto = alto / 12700 - 7.2
    for tamanio in range(maximo, minimo - 1, -1):
        por_renglon = max(1, int(util_ancho / (tamanio * 0.5) * 0.92))
        total = 0.0
        for i, (_, texto) in enumerate(b for b in bloques if b[1]):
            total += (12 if i else 0) + (tamanio + 1) * 1.2 + 4
            for linea in texto.split("\n"):
                total += -(-max(len(linea), 1) // por_renglon) * tamanio * 1.2 + 5
        if total <= util_alto:
            return tamanio
    return minimo


def _bloques(slide, x, y, ancho, alto, bloques, tamanio: int | None = None, color=AZUL, color_titulo=AZUL,
             maximo: int = 16, tarjeta=BLANCO):
    """Una caja con secciones: [(título, texto), ...]. Cada renglón del texto es un párrafo.

    Los títulos van en mayúsculas y en negrita, en el azul de GlobalAIze.
    Sin `tamanio`, usa el más grande que entra, hasta `maximo`: así la letra es lo más
    grande posible los días con poco texto, sin salirse de la caja los días con mucho.
    """
    if not any(texto for _, texto in bloques):
        return None  # sin nada que decir no queda una tarjeta vacía
    if tarjeta is not None:
        # La tarjeta ocupa la caja pedida y el texto va adentro, con margen
        _tarjeta_fondo(slide, x, y, ancho, alto, tarjeta)
        margen = Inches(0.15)
        x, y, ancho, alto = x + margen, y + margen, ancho - 2 * margen, alto - 2 * margen
    tamanio = tamanio or _tamanio_que_entra(bloques, ancho, alto, maximo)
    caja = slide.shapes.add_textbox(x, y, ancho, alto)
    marco = caja.text_frame
    marco.word_wrap = True
    primero = True
    for titulo, texto in bloques:
        if not texto:
            continue
        p = marco.paragraphs[0] if primero else marco.add_paragraph()
        if not primero:
            p.space_before = Pt(12)
        primero = False
        run = p.add_run()
        run.text = titulo.upper()
        run.font.size, run.font.bold, run.font.color.rgb = Pt(tamanio + 1), True, color_titulo
        p.space_after = Pt(4)
        for linea in texto.split("\n"):
            p = marco.add_paragraph()
            p.space_after = Pt(5)
            run = p.add_run()
            run.text = linea
            run.font.size, run.font.color.rgb = Pt(tamanio), color
    return caja


def _rectangulo(slide, x, y, ancho, alto, color, forma=1):
    rect = slide.shapes.add_shape(forma, x, y, ancho, alto)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def _sombra(forma, opacidad: float = 0.22) -> None:
    """Sombra suave hacia abajo, que despega la forma del fondo: python-pptx no la expone, va en el XML.

    _rectangulo ya dejó un <a:effectLst/> vacío (para no heredar la del tema); la sombra va adentro.
    """
    spPr = forma.fill._xPr
    efectos = spPr.find(qn("a:effectLst"))
    if efectos is None:
        efectos = spPr.makeelement(qn("a:effectLst"), {})
        spPr.append(efectos)
    sombra = efectos.makeelement(qn("a:outerShdw"), {
        "blurRad": "177800", "dist": "63500", "dir": "5400000", "algn": "t", "rotWithShape": "0",
    })
    color = sombra.makeelement(qn("a:srgbClr"), {"val": "081E40"})
    color.append(color.makeelement(qn("a:alpha"), {"val": str(int(opacidad * 100_000))}))
    sombra.append(color)
    efectos.append(sombra)


def _tarjeta_fondo(slide, x, y, ancho, alto, color=BLANCO):
    """Un recuadro de esquinas redondeadas, con borde fino y sombra: el fondo de cada gráfico o texto."""
    tarjeta = _rectangulo(slide, x, y, ancho, alto, color, forma=5)  # 5 = rectángulo redondeado
    tarjeta.adjustments[0] = 0.04
    if color == BLANCO:
        tarjeta.line.color.rgb = BORDE_TARJETA
        tarjeta.line.width = Pt(0.75)
    else:
        tarjeta.line.color.rgb = ACENTO
        tarjeta.line.width = Pt(1)
    _sombra(tarjeta, 0.22 if color == BLANCO else 0.35)
    return tarjeta


def _opacidad(forma, opacidad: float) -> None:
    """Relleno semitransparente: python-pptx no lo expone, va directo en el XML."""
    color = forma.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    color.append(color.makeelement(qn("a:alpha"), {"val": str(int(opacidad * 100_000))}))


def _degrade(forma, desde: RGBColor, hasta: RGBColor, angulo: float = 0) -> None:
    forma.fill.gradient()
    forma.fill.gradient_angle = angulo
    paradas = forma.fill.gradient_stops
    paradas[0].color.rgb, paradas[1].color.rgb = desde, hasta


def _imagen(slide, ruta: Path, x, y, ancho_max, alto_max, centrar: bool = True, tarjeta=BLANCO):
    """La imagen lo más grande posible dentro de la caja, sin deformarla, sobre su tarjeta con sombra.

    La tarjeta ocupa la caja entera, así todas las de una diapositiva arrancan y terminan
    a la misma altura; la imagen va centrada adentro, con margen. Con tarjeta=None va sola.
    """
    margen = Inches(0.12) if tarjeta is not None else 0
    if tarjeta is not None:
        _tarjeta_fondo(slide, x, y, ancho_max, alto_max, tarjeta)
    with Image.open(ruta) as im:
        w, h = im.size
    escala = min((ancho_max - 2 * margen) / w, (alto_max - 2 * margen) / h)
    ancho, alto = int(w * escala), int(h * escala)
    if centrar:
        x, y = x + (ancho_max - ancho) // 2, y + (alto_max - alto) // 2
    else:
        x, y = x + margen, y + margen
    return slide.shapes.add_picture(str(ruta), x, y, ancho, alto)


def _con_link(forma, url: str):
    forma.click_action.hyperlink.address = url
    return forma


def _links(slide, x, y, lado, color_rotulo=CLARO, separacion=None, omitir=()):
    """La fila de íconos con link (LinkedIn, GitHub, GlobalAIze), cada uno con su rótulo abajo.

    Cada ícono va centrado en su lugar de ancho `separacion`; con x=None, la fila entera
    va centrada en la diapositiva.
    """
    separacion = separacion or Inches(1.55)
    links = [link for link in LINKS if link[1] not in omitir]
    if x is None:
        x = (ANCHO - separacion * len(links)) // 2
    # El logo de GlobalAIze, con las esquinas redondeadas como los demás íconos
    redondeado = RECURSOS / "globalaize_logo_redondeado.png"
    ruta_globalaize = redondeado if redondeado.exists() else _icono_globalaize()
    for i, (icono, rotulo, url) in enumerate(links):
        ruta = ruta_globalaize if icono is None else RECURSOS / icono
        izquierda = x + i * separacion
        _con_link(slide.shapes.add_picture(str(ruta), izquierda + (separacion - lado) // 2, y, lado, lado), url)
        _texto(slide, izquierda, y + lado + Inches(0.03), separacion, Inches(0.5), rotulo, 11,
               color_rotulo, alinear=PP_ALIGN.CENTER, link=url)


def _encabezado(prs, titulo: str, fecha: date, numero: int, oscuro: bool = False, icono: Path | None = None,
                emoji: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NEGRO if oscuro else FONDO_SLIDE
    franja = _rectangulo(slide, 0, 0, ANCHO, Inches(0.9), OSCURO)
    if not oscuro:
        _degrade(franja, OSCURO, AZUL)
    _rectangulo(slide, 0, Inches(0.9), ANCHO, Inches(0.05), ACENTO)
    izquierda = Inches(0.45)
    if icono is not None:
        slide.shapes.add_picture(str(icono), Inches(0.45), Inches(0.2), Inches(0.5), Inches(0.5))
        izquierda = Inches(1.05)
    caja = _texto(slide, izquierda, Inches(0.14), Inches(9.5), Inches(0.6), titulo, 26, BLANCO, True)
    if emoji:
        # El emoji en un tramo aparte, con Segoe UI Emoji: con la tipografía del título
        # PowerPoint lo dibuja en blanco y negro o como un cuadrado
        parrafo = caja.text_frame.paragraphs[0]
        tramo = parrafo.runs[0]._r
        nuevo = parrafo.add_run()
        nuevo.text = f"{emoji} "
        nuevo.font.size, nuevo.font.name = Pt(24), "Segoe UI Emoji"
        tramo.addprevious(nuevo._r)
    _texto(slide, Inches(10.4), Inches(0.25), Inches(2.6), Inches(0.5), f"{fecha:%d/%m/%Y}", 16, BLANCO,
           alinear=PP_ALIGN.RIGHT)
    pie = CLARO if oscuro else GRIS
    _texto(slide, Inches(0.45), Inches(7.08), Inches(9), Inches(0.3), "Reporte Macroeconómico · Mauro E. Martinez · GlobalAIze",
           10, pie)
    _texto(slide, Inches(11.9), Inches(7.08), Inches(1), Inches(0.3), str(numero), 10, pie, alinear=PP_ALIGN.RIGHT)
    return slide


def _caja_ia(slide, x, y, ancho, alto, texto: str, rotulo: str = "Análisis IA del día"):
    """La tarjeta oscura del párrafo de IA, como en el mail."""
    _tarjeta_fondo(slide, x, y, ancho, alto, OSCURO)
    caja = _bloques(slide, x + Inches(0.2), y + Inches(0.08), ancho - Inches(0.35), alto - Inches(0.15),
                    [(f"🤖 {rotulo}", texto)], None, CLARO, ACENTO_CLARO, tarjeta=None)
    if caja is not None:
        caja.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    return caja


def _comentario(comentarios: dict, cid: str, palabra: str) -> str | None:
    """El comentario de Gemini de ese gráfico cuyo título menciona `palabra` (paralela, oficial, riesgo...)."""
    return next((texto for titulo, texto in comentarios.get(cid, []) if palabra in titulo.lower()), None)


def _variacion(hoy: float, ayer: float) -> tuple[str, RGBColor]:
    """Variación contra la rueda anterior: subir es rojo y bajar, verde (criterio macro del mail)."""
    var = (hoy / ayer - 1) * 100
    return f"{_signo(var)}% contra la rueda anterior", (ROJO if var > 0 else VERDE if var < 0 else GRIS)


def _tarjeta(slide, x, y, titulo: str, valor: str, detalle: str, color_detalle=GRIS):
    _tarjeta_fondo(slide, x, y, Inches(3.9), Inches(1.6))
    _rectangulo(slide, x, y + Inches(0.12), Inches(0.07), Inches(1.36), color_detalle if color_detalle != GRIS else AZUL)
    _texto(slide, x + Inches(0.25), y + Inches(0.1), Inches(3.5), Inches(0.4), titulo.upper(), 13, GRIS, True)
    _texto(slide, x + Inches(0.25), y + Inches(0.42), Inches(3.5), Inches(0.6), valor, 32, AZUL, True)
    _texto(slide, x + Inches(0.25), y + Inches(1.08), Inches(3.5), Inches(0.4), detalle, 14, color_detalle)


# ── Frases calculadas (cuando Gemini no comenta, o no hay comentario para ese gráfico) ──

def frase_paralelas(hoy, ayer) -> str:
    def dato(nombre, col):
        return f"{nombre} $ {_num(hoy[col])} ({_signo((hoy[col] / ayer[col] - 1) * 100)}%)"
    return (f"{dato('Blue', 'TCV_Blue')}, {dato('MEP', 'TCV_MEP')}, {dato('euro blue', 'TCV_Euro')} "
            f"y solidario $ {_num(hoy['Solidario'])}, siempre a la venta y contra la rueda anterior.")


def frase_oficiales(hoy, ayer) -> str:
    def dato(nombre, col):
        return f"{nombre} $ {_num(hoy[col])} ({_signo((hoy[col] / ayer[col] - 1) * 100)}%)"
    return (f"{dato('Billete BNA', 'TCV_Billete')} y {dato('divisas', 'TCV_Divisas')}, a la venta y contra la "
            "rueda anterior. El billete es el que compra la gente en el banco; divisas, el de las empresas.")


def frase_riesgo(df) -> str:
    ultimos = df["riesgo_pais"].head(charts.COTIZACIONES_A_MOSTRAR)
    hoy, ayer = df["riesgo_pais"].iloc[0], df["riesgo_pais"].iloc[1]
    return (f"{_num(hoy, 0)} puntos, {_signo(hoy - ayer, 0)} contra la rueda anterior. En las últimas "
            f"{len(ultimos)} ruedas estuvo entre {_num(ultimos.min(), 0)} y {_num(ultimos.max(), 0)}.\n"
            "Es lo que paga Argentina por encima de los bonos de Estados Unidos: cuanto más bajo, más barato le sale "
            "el crédito al país y a sus empresas.")


def frase_inflacion(inflacion: pd.DataFrame | None, df) -> str | None:
    """La última inflación mensual e interanual, y lo acumulado desde el inicio del gráfico de variaciones."""
    if inflacion is None or inflacion.empty:
        return None
    serie = inflacion.dropna(subset=["Inflación Mensual"])
    ultimo, previo = serie.iloc[-1], serie.iloc[-2] if len(serie) > 1 else None
    texto = f"La inflación de {indicadores.mes_y_anio(ultimo['Fecha'])} fue {_num(ultimo['Inflación Mensual'], 1)}%"
    if previo is not None:
        texto += f" ({indicadores.MESES[previo['Fecha'].month - 1]}: {_num(previo['Inflación Mensual'], 1)}%)"
    if pd.notna(ultimo.get("Inflación Anual")):
        texto += f" y la de los últimos 12 meses, {_num(ultimo['Inflación Anual'], 1)}%"
    texto += "."

    inicio = pd.Timestamp(charts.FECHA_INICIO_VARIACIONES)
    precios = ((serie[serie["Fecha"] >= inicio]["Inflación Mensual"] / 100 + 1).prod() - 1) * 100
    cotizaciones = df[["Fecha", "TCV_Blue", "Solidario"]].assign(Fecha=lambda d: pd.to_datetime(d["Fecha"].astype(str).str[:10]))
    antes = cotizaciones[cotizaciones["Fecha"] < inicio]
    if not antes.empty:
        base, hoy = antes.iloc[0], cotizaciones.iloc[0]
        blue = (hoy["TCV_Blue"] / base["TCV_Blue"] - 1) * 100
        solidario = (hoy["Solidario"] / base["Solidario"] - 1) * 100
        texto += (f"\nDesde {indicadores.mes_y_anio(inicio)}, los precios subieron {_num(precios, 1)}%, el dólar blue "
                  f"{_signo(blue, 1)}% y el solidario {_signo(solidario, 1)}%: el dólar "
                  f"{'subió menos que los precios, así que en pesos de hoy está más barato' if blue < precios else 'subió más que los precios, así que en pesos de hoy está más caro'}.")
    return texto


def frase_btc(btc: pd.DataFrame | None) -> str | None:
    if btc is None or btc.empty or "Close" not in btc:
        return None
    cierre = btc["Close"].dropna()
    if len(cierre) < 2:
        return None
    ultimo, fecha = float(cierre.iloc[-1]), cierre.index[-1]

    def cambio(dias: int) -> str | None:
        previos = cierre[cierre.index <= fecha - pd.Timedelta(days=dias)]
        if previos.empty and (fecha - cierre.index[0]).days >= dias - 10:
            previos = cierre.iloc[:1]  # la serie viene recortada a un año justo: su primer dato
        return None if previos.empty else f"{_signo((ultimo / float(previos.iloc[-1]) - 1) * 100, 1)}%"

    partes = [f"en {t} {c}" for t, c in (("7 días", cambio(7)), ("30 días", cambio(30)), ("el último año", cambio(365))) if c]
    texto = f"Bitcoin cerró en USD {_num(ultimo, 0)} el {fecha:%d/%m/%Y}"
    texto += (": " + ", ".join(partes) + ".") if partes else "."
    if "Volatility" in btc and pd.notna(btc["Volatility"].iloc[-1]):
        texto += f"\nVolatilidad diaria de las últimas 20 ruedas: {_num(float(btc['Volatility'].iloc[-1]), 1)}%."
    return texto


# ── Archivo ──────────────────────────────────────────────────────────────────

def _metadatos(prs, fecha: date) -> None:
    propiedades = prs.core_properties
    propiedades.title = f"Reporte Macroeconómico {fecha:%d/%m/%Y}"
    propiedades.author = AUTOR
    propiedades.last_modified_by = AUTOR
    propiedades.comments = ""
    propiedades.revision = 1
    # python-pptx escribe estas fechas como UTC
    ahora = datetime.now(UTC).replace(tzinfo=None, microsecond=0)
    propiedades.created = ahora
    propiedades.modified = ahora


def _colores_de_link(prs) -> None:
    """Los links en los azules de GlobalAIze: PowerPoint ignora el color del texto y usa el del tema."""
    import re

    tema = prs.slide_masters[0].part.part_related_by(RELATIONSHIP_TYPE.THEME)
    xml = tema.blob.decode("utf-8")
    for etiqueta, color in (("hlink", "1F4E79"), ("folHlink", "081E40")):
        xml = re.sub(rf"(<a:{etiqueta}>\s*<a:srgbClr val=\")[0-9A-Fa-f]{{6}}", rf"\g<1>{color}", xml)
    tema._blob = xml.encode("utf-8")


def _propiedades_de_aplicacion(prs) -> None:
    """docProps/app.xml con los datos de este archivo.

    La plantilla de python-pptx trae los de la PowerPoint en que se armó (Microsoft
    Macintosh PowerPoint 14, presentación 4:3, cero diapositivas). python-pptx no
    tiene API para esta parte, así que se reescribe entera; es un XML de propiedades
    opcionales y PowerPoint lo acepta así de corto.
    """
    parte = prs.part.package.part_related_by(RELATIONSHIP_TYPE.EXTENDED_PROPERTIES)
    parte._blob = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f"<Application>{AUTOR}</Application>"
        "<PresentationFormat>16:9</PresentationFormat>"
        f"<Slides>{len(prs.slides)}</Slides>"
        "</Properties>"
    ).encode()




def armar(
    df, imagenes: dict[str, Path], salida: Path, parrafo: str | None = None, comentarios: dict | None = None,
    explicaciones: dict | None = None, inflacion: pd.DataFrame | None = None, btc: pd.DataFrame | None = None,
) -> Path:
    """Arma la presentación en `salida`/ARCHIVO y devuelve su ruta.

    `df` es Fact_Mercado_Macro más nuevo primero (la fila 0 es el día del reporte),
    como lo devuelve data_access.leer_historico o como lo arma el pipeline.
    `imagenes` son los gráficos disponibles por nombre (charts.ORDEN_EN_MAIL); el que
    falte sale como "Gráfico no disponible".

    `parrafo` y `comentarios` son los textos de IA de esta corrida: el pipeline los
    pasa en memoria porque todavía no están en `df`. Sin ellos se leen de la fila 0
    (ai_paragraph y ai_secciones), que es lo que hace el script manual. `comentarios`
    tiene la forma de ia_generator.comentarios_por_grafico; {} significa "sin comentarios".
    `explicaciones` son las de indicadores.explicaciones; sin ellas, las diapositivas
    de agregados y deuda llevan solo lo ideal, sin el resumen con los últimos datos.
    `inflacion` (transformations.serie_inflacion) y `btc` (charts.preparar_btc) son
    para las frases de esos dos gráficos; sin ellos, esas diapositivas van sin frase.
    """
    hoy, ayer = df.iloc[0], df.iloc[1]
    fecha = _fecha(hoy["Fecha"])
    fwd_oficial, fwd_blue = transformations.forwards_fisher(df)
    resumen = email_report.resumen_ejecutivo(transformations.agregar_brechas_y_variaciones(df))
    if comentarios is None:
        comentarios = ia_generator.comentarios_por_grafico(hoy.get("ai_secciones"))
    if parrafo is None:
        guardado = hoy.get("ai_paragraph")
        parrafo = guardado if isinstance(guardado, str) and guardado else SIN_ANALISIS
    if explicaciones is None:
        explicaciones = indicadores.explicaciones({})

    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO
    _metadatos(prs, fecha)
    numero = iter(range(2, 100))  # la portada es la 1 y no lleva número

    # 1. Portada: blanca, con el logo de GlobalAIze y sus azules, los links y las tecnologías
    portada = prs.slides.add_slide(prs.slide_layouts[6])
    portada.background.fill.solid()
    portada.background.fill.fore_color.rgb = BLANCO
    for x, y, lado, opacidad in [(-1.6, -1.9, 4.6, 0.9), (10.9, -1.4, 3.8, 0.7), (11.6, 3.6, 2.6, 0.5)]:
        _opacidad(_rectangulo(portada, Inches(x), Inches(y), Inches(lado), Inches(lado), FONDO_TARJETA, forma=9), opacidad)
    lado_logo = Inches(1.45)
    # El logo grande es el link al sitio: por eso la fila de links de abajo no lo repite
    _con_link(portada.shapes.add_picture(str(_icono_globalaize()), (ANCHO - lado_logo) // 2, Inches(0.3), lado_logo, lado_logo),
              "https://www.globalaize.com")
    centrado = {"alinear": PP_ALIGN.CENTER}
    _texto(portada, 0, Inches(1.85), ANCHO, Inches(0.4), "ARGENTINA · INFORME DIARIO", 14, GRIS_LOGO, True, **centrado)
    _texto(portada, 0, Inches(2.2), ANCHO, Inches(1.0), "Reporte Macroeconómico", 50, OSCURO, True, **centrado)
    _texto(portada, 0, Inches(3.15), ANCHO, Inches(0.5),
           "Dólar · Riesgo país · Tasas · Inflación · Agregados monetarios · Deuda · Bitcoin", 18, GRIS_LOGO, **centrado)
    _texto(portada, 0, Inches(3.65), ANCHO, Inches(0.5), f"{fecha:%d/%m/%Y}", 24, AZUL, True, **centrado)
    _texto(portada, 0, Inches(4.25), ANCHO, Inches(0.4), "por Mauro E. Martinez · GlobalAIze", 15, GRIS_LOGO, **centrado)
    _links(portada, None, Inches(4.7), Inches(0.5), color_rotulo=OSCURO, separacion=Inches(1.8), omitir=("GlobalAIze",))

    # Pie: las tecnologías que intervienen, sobre una franja celeste
    _rectangulo(portada, 0, Inches(6.05), ANCHO, Inches(1.45), FONDO_TARJETA)
    _texto(portada, 0, Inches(6.1), ANCHO, Inches(0.3), "STACK", 10, GRIS_LOGO, True, **centrado)
    lugar, icono = Inches(0.78), Inches(0.42)
    inicio = (ANCHO - lugar * len(TECNOLOGIAS)) // 2
    for i, (archivo, rotulo) in enumerate(TECNOLOGIAS):
        x = inicio + i * lugar
        portada.shapes.add_picture(str(RECURSOS / "tecnologias" / f"{archivo}.png"), x + (lugar - icono) // 2,
                                   Inches(6.4), icono, icono)
        _texto(portada, x - Inches(0.05), Inches(6.9), lugar + Inches(0.1), Inches(0.3), rotulo, 9, OSCURO, **centrado)

    # 2. Tablero
    tablero = _encabezado(prs, "Tablero del día", fecha, next(numero), emoji="📋")
    tarjetas = [
        ("Dólar blue", f"$ {_num(hoy['TCV_Blue'])}", *_variacion(hoy["TCV_Blue"], ayer["TCV_Blue"])),
        ("Dólar MEP", f"$ {_num(hoy['TCV_MEP'])}", *_variacion(hoy["TCV_MEP"], ayer["TCV_MEP"])),
        ("Billete BNA", f"$ {_num(hoy['TCV_Billete'])}", *_variacion(hoy["TCV_Billete"], ayer["TCV_Billete"])),
        ("Riesgo país", f"{_num(hoy['riesgo_pais'], 0)} pts",
         f"{_signo(hoy['riesgo_pais'] - ayer['riesgo_pais'], 0)} pts contra la rueda anterior",
         ROJO if hoy["riesgo_pais"] > ayer["riesgo_pais"] else VERDE if hoy["riesgo_pais"] < ayer["riesgo_pais"] else GRIS),
        ("BADLAR (TEA)", f"{_num(hoy['bcra_tea'])}%", f"FED (EFFR): {_num(hoy['fed_tea'])}%", GRIS),
        ("Forward oficial 3 meses", f"$ {_num(fwd_oficial)}", f"Forward blue: $ {_num(fwd_blue)}", GRIS),
    ]
    for i, (titulo, valor, detalle, color) in enumerate(tarjetas):
        _tarjeta(tablero, Inches(0.45 + (i % 3) * 4.2), Inches(1.4 + (i // 3) * 2.0), titulo, valor, detalle, color)
    diferencia = abs(resumen["ahorro_valor"])
    sentido = "más barato" if resumen["ahorro_valor"] > 0 else "más caro"
    _tarjeta_fondo(tablero, Inches(0.45), Inches(5.35), Inches(12.4), Inches(1.4))
    _texto(tablero, Inches(0.65), Inches(5.45), Inches(12.0), Inches(1.2),
           f"Comprar USD {resumen['cantidad_usd']} al oficial con impuestos (solidario) sale $ {_num(diferencia)} "
           f"{sentido} que en el blue: el solidario está {_num(abs(resumen['brecha_solidario']))}% "
           f"{'por debajo' if resumen['brecha_solidario'] > 0 else 'por encima'} del blue.\n"
           "Verde: bajó contra la rueda anterior; rojo: subió. Desde la macro, que el dólar y el riesgo país bajen "
           "es la buena noticia.", 15, GRIS)

    # 3 y 4. Tipos de cambio y riesgo país, con los paneles por separado
    paneles = []
    with tempfile.TemporaryDirectory() as temporal:
        if charts.TIPOS_DE_CAMBIO in imagenes:
            try:
                paneles = charts.paneles_tipos_de_cambio(charts.preparar_datos(df), Path(temporal))
            except Exception:
                paneles = []  # sin paneles, va el gráfico entero como en el mail

        cambio = _encabezado(prs, "Tipos de cambio", fecha, next(numero), emoji="💵")
        _caja_ia(cambio, Inches(0.45), Inches(1.12), Inches(12.43), Inches(1.25), parrafo)
        cambio.notes_slide.notes_text_frame.text = parrafo
        if paneles:
            for j, (ruta, clave, frase) in enumerate([
                (paneles[0], "paralela", frase_paralelas(hoy, ayer)),
                (paneles[1], "oficial", frase_oficiales(hoy, ayer)),
            ]):
                x = Inches(0.45) + j * Inches(6.33)
                _imagen(cambio, ruta, x, Inches(2.5), Inches(6.1), Inches(2.35))
                texto = _comentario(comentarios, charts.cid(charts.TIPOS_DE_CAMBIO), clave)
                _bloques(cambio, x, Inches(4.95), Inches(6.1), Inches(2.0),
                         [("🤖 Análisis IA" if texto else "Resumen del día", texto or frase)], maximo=15)
        elif charts.TIPOS_DE_CAMBIO in imagenes:
            _imagen(cambio, imagenes[charts.TIPOS_DE_CAMBIO], Inches(0.45), Inches(2.5), Inches(12.43), Inches(4.45))
        else:
            _texto(cambio, Inches(0.6), Inches(4), Inches(12), Inches(1), "Gráfico no disponible para este día.", 18, GRIS)

        riesgo = _encabezado(prs, "Riesgo país", fecha, next(numero), emoji="🌎")
        if paneles:
            _imagen(riesgo, paneles[2], Inches(0.45), Inches(1.25), Inches(7.9), Inches(5.6))
        else:
            _texto(riesgo, Inches(0.6), Inches(3), Inches(8), Inches(1), "Gráfico no disponible para este día.", 18, GRIS)
        _bloques(riesgo, Inches(8.6), Inches(1.2), Inches(4.35), Inches(5.75),
                 [("Resumen del día", frase_riesgo(df)), ("🤖 Análisis IA", _comentario(comentarios, charts.cid(charts.TIPOS_DE_CAMBIO), "riesgo"))])

    # 5. Inflación
    inflacion_slide = _encabezado(prs, "Inflación y variaciones acumuladas", fecha, next(numero), emoji="📈")
    for j, nombre in enumerate([charts.INFLACION, charts.VARIACIONES]):
        x = Inches(0.45) + j * Inches(6.33)
        if nombre in imagenes:
            _imagen(inflacion_slide, imagenes[nombre], x, Inches(1.1), Inches(6.1), Inches(4.0))
        else:
            _texto(inflacion_slide, x, Inches(3), Inches(6), Inches(1), "Gráfico no disponible para este día.", 18, GRIS)
    frase = frase_inflacion(inflacion, df)
    if frase:
        _bloques(inflacion_slide, Inches(0.45), Inches(5.2), Inches(12.43), Inches(1.85), [("Resumen del día", frase)])

    # 6 y 7. Agregados y deuda: el resumen del día y lo ideal; el texto completo, en las notas
    for titulo, emoji, nombre, cid in [("Agregados monetarios", "💰", charts.AGREGADOS, indicadores.CID_AGREGADOS),
                                       ("Endeudamiento, en dólares", "💸", charts.DEUDA, indicadores.CID_DEUDA)]:
        slide = _encabezado(prs, titulo, fecha, next(numero), emoji=emoji)
        if nombre in imagenes:
            _imagen(slide, imagenes[nombre], Inches(0.45), Inches(1.1), Inches(6.7), Inches(5.9))
        else:
            _texto(slide, Inches(0.6), Inches(3), Inches(7), Inches(1), "Gráfico no disponible para este día.", 18, GRIS)
        explicacion = explicaciones.get(cid)
        if explicacion:
            parrafos = explicacion["texto"].split("\n")
            ideal = parrafos[-1].removeprefix("¿Qué sería lo ideal?").strip()
            _bloques(slide, Inches(7.35), Inches(1.15), Inches(5.6), Inches(5.85),
                     [("Resumen del día", explicacion["dato"]), ("¿Qué sería lo ideal?", ideal)])
            slide.notes_slide.notes_text_frame.text = "\n\n".join(filter(None, [*parrafos, explicacion["dato"]]))

    # 8. Bitcoin, en negro como su gráfico
    bitcoin = _encabezado(prs, "Bitcoin", fecha, next(numero), oscuro=True, icono=RECURSOS / "bitcoin.png")
    if charts.BTC in imagenes:
        _imagen(bitcoin, imagenes[charts.BTC], Inches(0.45), Inches(1.15), Inches(7.9), Inches(5.8), tarjeta=TARJETA_OSCURA)
    else:
        _texto(bitcoin, Inches(0.6), Inches(3), Inches(8), Inches(1), "Gráfico no disponible para este día.", 18, CLARO)
    _bloques(bitcoin, Inches(8.6), Inches(1.2), Inches(4.35), Inches(5.75),
             [("Resumen del día", frase_btc(btc)), ("🤖 Análisis IA", _comentario(comentarios, charts.cid(charts.BTC), "bitcoin"))], None, CLARO, ACENTO_CLARO, tarjeta=TARJETA_OSCURA)

    # 9. Fuentes y contacto
    cierre = _encabezado(prs, "Fuentes y contacto", fecha, next(numero), emoji="📚")
    _tarjeta_fondo(cierre, Inches(0.35), Inches(1.1), Inches(7.85), Inches(5.55))
    _texto(cierre, Inches(0.55), Inches(1.25), Inches(7.5), Inches(0.4), "FUENTES CONSULTADAS", 14, AZUL, True)
    caja = cierre.shapes.add_textbox(Inches(0.55), Inches(1.7), Inches(7.5), Inches(4.9))
    marco = caja.text_frame
    marco.word_wrap = True
    for i, (nombre, url) in enumerate(FUENTES):
        p = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        p.space_after = Pt(9)
        run = p.add_run()
        run.text = f"{nombre}: "
        run.font.size, run.font.bold, run.font.color.rgb = Pt(15), True, AZUL
        run = p.add_run()
        run.text = url
        run.hyperlink.address = url
        run.font.size, run.font.color.rgb = Pt(14), CELESTE
    _tarjeta_fondo(cierre, Inches(8.45), Inches(1.1), Inches(4.43), Inches(5.55), OSCURO)
    _texto(cierre, Inches(8.75), Inches(1.4), Inches(4), Inches(0.4), "CONTACTO", 14, ACENTO_CLARO, True)
    _texto(cierre, Inches(8.75), Inches(1.85), Inches(4), Inches(0.9),
           "Reporte hecho íntegramente en Python por Mauro E. Martinez, de GlobalAIze (sitio en construcción).", 14, CLARO)
    _links(cierre, Inches(8.65), Inches(3.0), Inches(0.5), separacion=Inches(1.03))
    _texto(cierre, Inches(8.75), Inches(4.45), Inches(4), Inches(0.8),
           f"☕ ¿Te sirvió el reporte? Podés apoyar el proyecto al alias {ALIAS}", 15, BLANCO, True)
    _texto(cierre, Inches(8.75), Inches(5.4), Inches(4), Inches(0.9),
           f"📧 Consultas, desuscripciones o propuestas:\n{MAIL}", 14, CLARO, link=f"mailto:{MAIL}")
    _texto(cierre, 0, Inches(6.72), ANCHO, Inches(0.35),
           "Análisis generado con Gemini a partir de los datos del día. No es asesoramiento financiero.", 10, GRIS,
           alinear=PP_ALIGN.CENTER)

    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)
    _colores_de_link(prs)
    _propiedades_de_aplicacion(prs)
    ruta = salida / ARCHIVO
    # Primero a un temporal y después se reemplaza: si el guardado falla a mitad de
    # camino, queda entera la versión anterior y no un archivo roto. El temporal
    # termina en .pptx, así que git lo ignora igual que al definitivo.
    temporal = ruta.with_name(f"~tmp {ARCHIVO}")
    try:
        prs.save(temporal)
        os.replace(temporal, ruta)
    finally:
        temporal.unlink(missing_ok=True)
    return ruta
