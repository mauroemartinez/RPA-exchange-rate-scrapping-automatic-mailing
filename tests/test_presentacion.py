"""Presentación ejecutiva: arma el .pptx del día con datos sintéticos, sin red."""

import zipfile
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

import presentacion_ejecutiva
from conftest import jpeg_minimo
from reporte import charts, indicadores, presentacion


def _textos(slide) -> str:
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


@pytest.fixture
def imagenes(tmp_path):
    rutas = {}
    for nombre in charts.ORDEN_EN_MAIL:
        rutas[nombre] = tmp_path / nombre
        rutas[nombre].write_bytes(jpeg_minimo())
    return rutas


def test_arma_las_nueve_diapositivas(historico, imagenes, tmp_path):
    ruta = presentacion.armar(historico, imagenes, tmp_path / "salida")
    prs = Presentation(ruta)

    assert len(prs.slides) == 9
    assert ruta == tmp_path / "salida" / presentacion.ARCHIVO
    tablero = _textos(prs.slides[1])
    assert "DÓLAR BLUE" in tablero and "RIESGO PAÍS" in tablero and "FORWARD OFICIAL 3 MESES" in tablero
    assert historico["ai_paragraph"].iloc[0] in _textos(prs.slides[2])


def test_el_nombre_es_fijo_y_se_pisa_cada_dia(historico, imagenes, tmp_path):
    primera = presentacion.armar(historico, imagenes, tmp_path)
    segunda = presentacion.armar(historico.iloc[1:].reset_index(drop=True), imagenes, tmp_path)
    assert primera == segunda and [p.name for p in tmp_path.glob("*.pptx")] == [presentacion.ARCHIVO]


def test_los_metadatos_no_le_dan_credito_a_nadie_mas(historico, imagenes, tmp_path):
    ruta = presentacion.armar(historico, imagenes, tmp_path)
    propiedades = Presentation(ruta).core_properties
    fecha = date.fromisoformat(historico["Fecha"].iloc[0])

    assert propiedades.title == f"Reporte Macroeconómico {fecha:%d/%m/%Y}"
    assert propiedades.author == propiedades.last_modified_by == presentacion.AUTOR
    assert propiedades.comments == "" and propiedades.created.year >= 2026
    # La plantilla de python-pptx trae el nombre de su autor y un comentario propio
    with zipfile.ZipFile(ruta) as pptx:
        xml = "".join(pptx.read(n).decode("utf-8", "replace") for n in pptx.namelist() if n.endswith((".xml", ".rels")))
    for ajeno in ("Steve Canny", "python-pptx"):
        assert ajeno.lower() not in xml.lower()


def test_con_comentarios_por_grafico_de_la_fila(historico, imagenes, tmp_path):
    historico = historico.copy()
    historico["ai_secciones"] = None
    historico.at[0, "ai_secciones"] = {"paralelas": "Comentario de paralelas.", "oficiales": "Comentario de oficiales.",
                                       "riesgo_pais": "Comentario de riesgo.", "btc": "Comentario de BTC."}
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    assert "Comentario de paralelas." in _textos(prs.slides[2])
    assert "Comentario de BTC." in _textos(prs.slides[7])


def test_los_textos_de_la_corrida_le_ganan_a_la_fila(historico, imagenes, tmp_path):
    """El pipeline pasa el párrafo y los comentarios en memoria: todavía no están en la fila."""
    historico = historico.copy()
    historico["ai_secciones"] = None
    historico.at[0, "ai_secciones"] = {"paralelas": "Comentario viejo de la fila."}

    prs = Presentation(presentacion.armar(
        historico, imagenes, tmp_path, parrafo="Resumen de esta corrida.",
        comentarios={"image1": [("Paralelas", "Comentario en memoria.")]},
    ))
    assert "Resumen de esta corrida." in _textos(prs.slides[2])
    assert "Comentario en memoria." in _textos(prs.slides[2])
    assert "Comentario viejo de la fila." not in _textos(prs.slides[2])

    # {} es "sin comentarios", no "buscalos en la fila"
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path, comentarios={}))
    assert "Comentario viejo de la fila." not in _textos(prs.slides[2])


def test_la_fila_recien_armada_trae_la_fecha_como_date(historico, imagenes, tmp_path):
    historico = historico.copy()
    historico["Fecha"] = historico["Fecha"].astype(object)
    historico.at[0, "Fecha"] = date.fromisoformat(historico["Fecha"].iloc[0])
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    assert historico["Fecha"].iloc[0].strftime("%d/%m/%Y") in _textos(prs.slides[0])


def test_sin_parrafo_lo_dice(historico, imagenes, tmp_path):
    historico = historico.copy()
    historico.loc[0, "ai_paragraph"] = None
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    assert presentacion.SIN_ANALISIS in _textos(prs.slides[2])


def test_sin_un_grafico_lo_avisa(historico, imagenes, tmp_path):
    del imagenes[charts.BTC]
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    assert "Gráfico no disponible" in _textos(prs.slides[7])


def test_el_script_manual_usa_el_mismo_armado(historico, imagenes, tmp_path, monkeypatch):
    class Engine:
        def dispose(self):
            pass

    monkeypatch.setattr(presentacion_ejecutiva.data_access, "crear_engine", Engine)
    monkeypatch.setattr(presentacion_ejecutiva.data_access, "leer_historico",
                        lambda engine, respaldo_csv=True: (historico.copy(), "supabase"))

    ruta = presentacion_ejecutiva.main(["--salida", str(tmp_path / "manual"), "--graficos", str(tmp_path)])

    assert ruta == tmp_path / "manual" / presentacion.ARCHIVO
    prs = Presentation(ruta)
    assert len(prs.slides) == 9 and historico["ai_paragraph"].iloc[0] in _textos(prs.slides[2])


def test_agregados_y_deuda_llevan_su_explicacion(historico, imagenes, tmp_path, series_indicadores):
    # Sin explicaciones (el script manual): el texto fijo, sin la frase con datos
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    # Sin la frase del día, igual va lo ideal, con su título
    assert "¿QUÉ SERÍA LO IDEAL?" in _textos(prs.slides[5])
    # En la diapositiva, "lo ideal" (el último párrafo); el texto completo, en las notas
    ideal = indicadores.TEXTO_DEUDA.split("\n")[-1].removeprefix("¿Qué sería lo ideal?").strip()
    assert ideal in _textos(prs.slides[6])
    notas = prs.slides[6].notes_slide.notes_text_frame.text
    assert all(parrafo in notas for parrafo in indicadores.TEXTO_DEUDA.split("\n"))

    # Con las de la corrida, también la frase con los últimos datos
    series, provisorios = series_indicadores
    explicaciones = indicadores.explicaciones(series, provisorios)
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path, explicaciones=explicaciones))
    assert explicaciones[indicadores.CID_DEUDA]["dato"] in _textos(prs.slides[6])
    assert explicaciones[indicadores.CID_AGREGADOS]["dato"] in _textos(prs.slides[5])


def test_los_metadatos_de_aplicacion_son_los_de_este_archivo(historico, imagenes, tmp_path):
    ruta = presentacion.armar(historico, imagenes, tmp_path)
    with zipfile.ZipFile(ruta) as pptx:
        app = pptx.read("docProps/app.xml").decode("utf-8")
    assert "Macintosh" not in app and "4:3" not in app
    assert f"<Slides>{len(Presentation(ruta).slides)}</Slides>" in app


def test_si_el_guardado_falla_queda_la_version_anterior(historico, imagenes, tmp_path, monkeypatch):
    ruta = presentacion.armar(historico, imagenes, tmp_path)
    anterior = ruta.read_bytes()

    def falla(self, destino):
        Path(destino).write_bytes(b"a medio escribir")
        raise OSError("disco lleno")

    monkeypatch.setattr(presentacion.Presentation().__class__, "save", falla)
    with pytest.raises(OSError):
        presentacion.armar(historico, imagenes, tmp_path)

    assert ruta.read_bytes() == anterior
    assert [p.name for p in tmp_path.glob("*.pptx")] == [presentacion.ARCHIVO]  # sin el temporal


def test_portada_con_bitcoin_y_links(historico, imagenes, tmp_path):
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    portada = prs.slides[0]
    assert "Bitcoin" in _textos(portada)
    links = {s.click_action.hyperlink.address for s in portada.shapes if s.click_action.hyperlink.address}
    assert {url for _, _, url in presentacion.LINKS} <= links


def test_cierre_con_fuentes_y_alias(historico, imagenes, tmp_path):
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    cierre = _textos(prs.slides[8])
    assert presentacion.ALIAS in cierre and "Fuentes y contacto" in cierre
    assert all(url in cierre for _, url in presentacion.FUENTES)


def test_subir_es_rojo_y_bajar_verde():
    """Desde la macro, que el dólar baje es la buena noticia."""
    assert presentacion._variacion(110, 100)[1] == presentacion.ROJO
    assert presentacion._variacion(90, 100)[1] == presentacion.VERDE


def test_bitcoin_en_negro_con_su_moneda(historico, imagenes, tmp_path):
    prs = Presentation(presentacion.armar(historico, imagenes, tmp_path))
    bitcoin = prs.slides[7]
    assert bitcoin.background.fill.fore_color.rgb == presentacion.NEGRO
    assert any(s.shape_type == 13 for s in bitcoin.shapes)  # 13 = imagen: la moneda del título


def test_frases_de_inflacion_y_btc(historico):
    import pandas as pd

    from reporte import transformations

    inflacion = transformations.serie_inflacion([(date(2025, m, 1), 2.0) for m in range(1, 13)]
                                                + [(date(2026, m, 1), 1.5) for m in range(1, 9)])
    frase = presentacion.frase_inflacion(inflacion, historico)
    assert frase.startswith("La inflación de agosto de 2026 fue 1,5% (julio: 1,5%)")
    fechas = pd.date_range("2025-10-08", "2026-10-08", freq="D")
    btc = pd.DataFrame({"Close": range(1, len(fechas) + 1)}, index=fechas, dtype=float)
    assert "Bitcoin cerró en USD 366 el 08/10/2026" in presentacion.frase_btc(btc)
    assert "el último año" in presentacion.frase_btc(btc)
    assert presentacion.frase_btc(None) is None and presentacion.frase_inflacion(None, historico) is None
