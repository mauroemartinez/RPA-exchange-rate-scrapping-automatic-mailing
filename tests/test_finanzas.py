"""Deuda bruta de la Secretaría de Finanzas: el link de la página, la lectura del Excel y sus controles."""

import io
from datetime import date, datetime

import httpx
import openpyxl
import pytest

from reporte.scrapers import agregados, finanzas
from reporte.scrapers.utils import ScraperError

PAGINA_HTML = """
<html><body>
<p>Últimos datos actualizados sobre el estado de la deuda pública bruta.</p>
<table class="table"><tbody>
<tr><td>Informe mensual</td><td><a href="/sites/default/files/informe.pdf">Descargar</a></td></tr>
<tr> <td data-label=&nbsp;>Serie mensual 2019 - Agosto 2026</td>
     <td><a href="/sites/default/files/boletin_mensual_31_08_2026_1.xlsx" class="btn">Descargar</a></td> </tr>
</tbody></table>
</body></html>
"""


def _planilla(meses: int = 30, unidad: str = "Año 2019 / 2026 - Datos en millones de U$S",
              rotulo: str = "A- DEUDA BRUTA ( I + II  + III)", valor=lambda k: 400_000 + 1_000 * k) -> bytes:
    """Un Excel con la forma del de la Secretaría: índice, título, unidad, meses y el total.

    Los dos últimos meses van como texto con asterisco, como los provisorios del original.
    """
    libro = openpyxl.Workbook()
    indice = libro.active
    indice.title = "Indice"
    indice["B2"] = "DEUDA BRUTA DE LA ADMINISTRACIÓN CENTRAL"
    hoja = libro.create_sheet("A.1")
    hoja["B5"] = "DEUDA BRUTA DE LA ADMINISTRACIÓN CENTRAL (**)"
    hoja["B6"] = unidad
    for k in range(meses):
        anio, mes = 2024 + (2 + k) // 12, (2 + k) % 12 + 1  # desde marzo de 2024
        columna = 3 + k
        if k >= meses - 2:
            abrev = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"][mes - 1]
            hoja.cell(row=9, column=columna, value=f"{abrev}-{anio % 100:02d} (*)")
        else:
            hoja.cell(row=9, column=columna, value=datetime(anio, mes, 1))
        hoja.cell(row=10, column=columna, value=valor(k))
        hoja.cell(row=12, column=columna, value=valor(k) - 2_000)
    hoja.cell(row=9, column=3 + meses, value="en %")
    hoja.cell(row=10, column=3 + meses, value=1)
    hoja["B10"] = rotulo
    hoja["B12"] = "B- DEUDA BRUTA (EXCLUIDA LA ELEGIBLE PENDIENTE DE REESTRUCTURACIÓN) ( I + II )"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_lee_la_fila_del_total_por_su_rotulo_con_los_provisorios():
    puntos, provisorios = finanzas.leer_planilla(_planilla())
    assert len(puntos) == 30
    assert puntos[0] == (date(2024, 3, 31), 400_000.0)
    # Los meses van al último día: son saldos a fin de mes
    assert puntos[-1] == (date(2026, 8, 31), 429_000.0)
    assert provisorios == {date(2026, 7, 31), date(2026, 8, 31)}
    # La columna "en %" y la fila B quedan afuera
    assert all(v >= 400_000 for _, v in puntos)


def test_el_rotulo_puede_venir_con_otros_espacios_y_sin_tildes():
    puntos, _ = finanzas.leer_planilla(_planilla(rotulo="A -  Deuda Bruta (I + II + III)"))
    assert len(puntos) == 30


@pytest.mark.parametrize("cambios, mensaje", [
    ({"unidad": "Datos en miles de millones de pesos"}, "millones de U\\$S"),
    ({"rotulo": "TOTAL GENERAL"}, "A- DEUDA BRUTA"),
    ({"meses": 12}, "meses"),
    ({"valor": lambda k: 400 + k}, "fuera de rango"),
])
def test_si_la_planilla_cambia_levanta_en_vez_de_leer_mal(cambios, mensaje):
    with pytest.raises(ValueError, match=mensaje):
        finanzas.leer_planilla(_planilla(**cambios))


def test_el_link_sale_de_la_fila_de_la_serie_mensual():
    url = finanzas.url_planilla(PAGINA_HTML, "https://www.argentina.gob.ar/economia/finanzas/datos-mensuales")
    assert url == "https://www.argentina.gob.ar/sites/default/files/boletin_mensual_31_08_2026_1.xlsx"


def test_sin_la_fila_de_la_serie_mensual_levanta():
    with pytest.raises(ValueError, match="Serie mensual"):
        finanzas.url_planilla("<table><tr><td>Otra cosa</td><td><a href='x.pdf'>x</a></td></tr></table>")


def test_descarga_de_punta_a_punta(monkeypatch):
    pedidos = []

    def responder(request: httpx.Request) -> httpx.Response:
        pedidos.append(str(request.url))
        if request.url.path.endswith(".xlsx"):
            return httpx.Response(200, content=_planilla())
        return httpx.Response(200, text=PAGINA_HTML)

    with httpx.Client(transport=httpx.MockTransport(responder), follow_redirects=True) as client:
        puntos, provisorios = finanzas.descargar(client)

    assert pedidos == [finanzas.PAGINA, "https://www.argentina.gob.ar/sites/default/files/boletin_mensual_31_08_2026_1.xlsx"]
    assert len(puntos) == 30 and len(provisorios) == 2


def test_un_error_de_la_pagina_es_un_scraper_error():
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(404))) as client:
        with pytest.raises(ScraperError, match="Secretaría de Finanzas"):
            finanzas.descargar(client)


def test_la_serie_se_describe_en_el_catalogo():
    serie = agregados.CATALOGO["deuda_bruta_tesoro"]
    assert serie is agregados.DEUDA_BRUTA
    assert (serie.frecuencia, serie.unidad, serie.fuente) == ("M", "millones de USD", "Secretaría de Finanzas")
    assert serie.id_fuente == "A.1, A- DEUDA BRUTA"
    assert agregados.POR_CLAVE["base_monetaria"].id_fuente == "15"


def test_sin_openpyxl_el_modulo_se_importa_igual(monkeypatch):
    """pipeline.py importa este módulo al arrancar: sin openpyxl, el mail tiene que salir igual."""
    import importlib
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    modulo = importlib.reload(finanzas)
    with pytest.raises(ImportError):
        modulo.leer_planilla(b"cualquier cosa")
