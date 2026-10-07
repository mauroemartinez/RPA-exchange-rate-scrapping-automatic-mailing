"""Orquestación de pipeline.correr() con todos los efectos externos reemplazados."""

import contextlib
import smtplib
import subprocess
from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path

import pytest

import charts
import ia_generator
import pipeline
import presentacion
from conftest import HOY
from scrapers import agregados
from scrapers.utils import ScraperError


class EngineFalso:
    def dispose(self):
        pass


@pytest.fixture
def entorno(historico, resultados, btc_crudo, tmp_path, monkeypatch):
    """(Dependencias falsas, registro de lo que se hizo, carpeta de salida).

    PREVIEWS apunta a una carpeta temporal: un test que no pase `salida` escribe
    ahí y no en el Previews/ del repo, que se commitea y se pushea solo.
    """
    monkeypatch.setattr(pipeline, "PREVIEWS", tmp_path / "Previews")
    hechos = {"filas": [], "parrafos": [], "mails": [], "alertas": [], "previews": [], "limpiezas": [],
              "presentaciones": []}

    def guardar_fila(engine, fila, sobrescribir=False):
        hechos["filas"].append((fila, sobrescribir))
        return True

    def generar_parrafo(engine, fecha_esperada=None):
        hechos["parrafos"].append(fecha_esperada)
        return "Párrafo de Gemini"

    deps = pipeline.Dependencias(
        crear_engine=EngineFalso,
        leer_historico=lambda engine: (historico.copy(), "supabase"),
        guardar_fila=guardar_fila,
        candado=lambda engine: contextlib.nullcontext(True),
        scrapear=lambda: (resultados.bna, resultados.dolarhoy, resultados.ambito,
                          resultados.riesgo_pais, resultados.bcra, resultados.fed),
        descargar_btc=lambda desde, hasta: btc_crudo.copy(),
        generar_parrafo=generar_parrafo,
        enviar_mail=lambda mensaje, destinatarios: hechos["mails"].append((mensaje, destinatarios)),
        actualizar_previews=lambda repo, archivos=None: hechos["previews"].append(repo) or (True, "ok"),
        publicar_presentacion=lambda repo, archivo, mensaje=None: hechos["presentaciones"].append((archivo, mensaje))
        or (True, "publicada"),
        alertar=lambda asunto, cuerpo: hechos["alertas"].append(asunto) or True,
        alertar_scraper=lambda exc: hechos["alertas"].append("scraper") or True,
        alertar_validacion=lambda exc: hechos["alertas"].append("validacion") or True,
        tabla_series=lambda engine: False,
        columna_secciones=lambda engine: False,
        limpiar_secciones=lambda engine, fecha: hechos["limpiezas"].append(fecha) or 1,
        feriado=lambda fecha: None,
        hoy=lambda: HOY,
        ahora=lambda: datetime(2026, 10, 6, 16, 43),
    )
    return deps, hechos, tmp_path


def _estados(resultado):
    return {e.nombre: e.estado for e in resultado.etapas}


def test_corrida_completa(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert r.estado == "ok" and r.exitosa
    assert _estados(r) == {
        "historico": "ok", "scraping": "ok", "validacion": "ok", "persistencia": "ok",
        "ia": "ok", "graficos": "ok", "mail": "ok", "presentacion": "ok", "previews": "omitida", "series": "omitida",
    }
    fila, sobrescribir = hechos["filas"][0]
    assert fila["Fecha"].iloc[0] == HOY and sobrescribir is False
    assert hechos["parrafos"] == [HOY]
    assert len(hechos["mails"]) == 2
    assert {p.name for p in salida.glob("*.jpg")} == set(charts.ORDEN_EN_MAIL)
    assert (salida / presentacion.ARCHIVO).exists()
    assert hechos["alertas"] == []
    html = hechos["mails"][0][0].get_payload()[0].get_payload(decode=True).decode()
    assert "Párrafo de Gemini" in html


def test_scraper_caido_alerta_y_corta(entorno):
    deps, hechos, salida = entorno

    def scrapear():
        raise ScraperError("BNA", "leer cotizaciones", TimeoutError("timeout"))

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, scrapear=scrapear))

    assert r.estado == "error" and not r.exitosa
    assert _estados(r)["scraping"] == "error"
    assert "persistencia" not in _estados(r)
    assert hechos["alertas"] == ["scraper"]  # sin alerta de resumen duplicada
    assert hechos["filas"] == [] and hechos["mails"] == []


def test_fila_invalida_alerta_y_no_persiste(entorno, resultados):
    deps, hechos, salida = entorno
    rota = (resultados.bna, {"TCC_Blue": 0.0, "TCV_Blue": 1555.0}, resultados.ambito,
            resultados.riesgo_pais, resultados.bcra, resultados.fed)

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, scrapear=lambda: rota))

    assert _estados(r)["validacion"] == "error"
    assert hechos["alertas"] == ["validacion"]
    assert hechos["filas"] == [] and hechos["mails"] == []


def test_si_la_fila_de_hoy_ya_existe_no_se_repite(entorno, historico):
    deps, hechos, salida = entorno
    con_hoy = historico.copy()
    con_hoy.loc[0, "Fecha"] = str(HOY)

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, leer_historico=lambda e: (con_hoy, "supabase")))

    assert r.estado == "omitida" and r.exitosa
    assert hechos["filas"] == [] and hechos["mails"] == []


def test_forzar_rehace_el_dia_pisando_la_fila(entorno, historico):
    deps, hechos, salida = entorno
    con_hoy = historico.copy()
    con_hoy.loc[0, "Fecha"] = str(HOY)

    r = pipeline.correr(
        pipeline.Opciones(salida=salida, forzar=True), replace(deps, leer_historico=lambda e: (con_hoy, "supabase"))
    )

    assert r.estado == "ok"
    assert hechos["filas"][0][1] is True  # sobrescribir
    assert len(hechos["mails"]) == 2


def test_otra_corrida_en_curso(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, candado=lambda e: contextlib.nullcontext(False)))
    assert r.estado == "omitida"
    assert _estados(r) == {"control": "omitida"}


def test_mail_fallido_pone_la_corrida_en_rojo_y_alerta(entorno):
    deps, hechos, salida = entorno

    def enviar(mensaje, destinatarios):
        raise smtplib.SMTPAuthenticationError(535, b"Username and Password not accepted")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, enviar_mail=enviar))

    assert r.estado == "error" and not r.exitosa
    assert _estados(r)["mail"] == "error"
    assert "535" in next(e.detalle for e in r.etapas if e.nombre == "mail")
    assert len(hechos["alertas"]) == 1 and "errores" in hechos["alertas"][0]


def test_sin_btc_el_mail_sale_igual_sin_ese_grafico(entorno):
    deps, hechos, salida = entorno

    def descargar(desde, hasta):
        raise ScraperError("Yahoo Finance", "descargar BTC-USD", ValueError("respuesta vacía"))

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, descargar_btc=descargar))

    assert r.estado == "advertencia" and r.exitosa
    assert _estados(r)["graficos"] == "advertencia"
    mensaje = hechos["mails"][0][0]
    cids = [p.get("Content-ID") for p in mensaje.walk() if p.get_content_type() == "image/jpeg"]
    assert cids == ["<image1>", "<image2>", "<image3>"]
    assert "cid:image4" not in mensaje.get_payload()[0].get_payload(decode=True).decode()


def test_si_no_se_guarda_la_fila_no_se_llama_a_gemini(entorno):
    deps, hechos, salida = entorno

    def guardar_fila(engine, fila, sobrescribir=False):
        raise ConnectionError("Supabase no responde")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, guardar_fila=guardar_fila))

    assert _estados(r)["persistencia"] == "error"
    assert _estados(r)["ia"] == "omitida"
    assert hechos["parrafos"] == []
    assert len(hechos["mails"]) == 2  # el reporte sale igual, con el mensaje de reemplazo
    html = hechos["mails"][0][0].get_payload()[0].get_payload(decode=True).decode()
    assert ia_generator.MENSAJE_FALLA in html
    assert r.estado == "error"


def test_gemini_sin_parrafo_es_advertencia(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, generar_parrafo=lambda e, fecha_esperada: ""))
    assert _estados(r)["ia"] == "advertencia"
    assert r.estado == "advertencia"


def test_dry_run_no_escribe_ni_manda_nada(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), deps)

    assert r.estado == "ok"
    assert _estados(r)["persistencia"] == "omitida"
    assert _estados(r)["ia"] == "omitida"
    assert _estados(r)["mail"] == "omitida"
    assert hechos["filas"] == [] and hechos["parrafos"] == [] and hechos["mails"] == [] and hechos["previews"] == []
    assert (salida / "mail.eml").exists() and (salida / "mail_preview.html").exists()


def test_dry_run_con_enviar_a_manda_solo_a_esa_direccion(entorno):
    deps, hechos, salida = entorno
    pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida, enviar_a=["yo@example.com"]), deps)

    assert len(hechos["mails"]) == 1
    _, destinatarios = hechos["mails"][0]
    assert destinatarios == ["yo@example.com", "remitente@example.com"]
    assert hechos["filas"] == []


def test_dry_run_no_manda_alertas(entorno):
    deps, hechos, salida = entorno

    def scrapear():
        raise ScraperError("BNA", "leer cotizaciones", TimeoutError("timeout"))

    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), replace(deps, scrapear=scrapear))
    assert r.estado == "error"
    assert hechos["alertas"] == []


def test_resultado_serializable(entorno):
    import json

    deps, _, salida = entorno
    datos = json.loads(json.dumps(pipeline.correr(pipeline.Opciones(salida=salida), deps).como_dict()))
    assert datos["fecha"] == "2026-10-06" and datos["exitosa"] is True and datos["estado"] == "ok"
    assert {"nombre", "estado", "segundos", "detalle"} <= set(datos["etapas"][0])


def test_cli_valida_las_direcciones_de_enviar_a():
    with pytest.raises(SystemExit):
        pipeline.main(["--dry-run", "--enviar-a", "no-es-un-mail"])


def test_cli_devuelve_1_si_la_corrida_falla(monkeypatch):
    fallida = pipeline.ResultadoCorrida(fecha=HOY, estado="error")
    monkeypatch.setattr(pipeline, "correr", lambda opciones: fallida)
    monkeypatch.setattr(pipeline, "configurar_logging", lambda archivo: None)
    assert pipeline.main(["--dry-run"]) == 1


def test_el_csv_adjunto_sale_de_supabase_con_el_parrafo_de_hoy(entorno, historico):
    deps, hechos, salida = entorno
    pipeline.correr(pipeline.Opciones(salida=salida), deps)

    con_csv = next(m for m, d in hechos["mails"] if "csv@example.com" in d)
    adjunto = next(p for p in con_csv.walk() if p.get_content_type() == "text/csv")
    lineas = adjunto.get_payload(decode=True).decode("utf-8").splitlines()
    assert lineas[0].startswith("Fecha,TCC_Blue,TCV_Blue") and lineas[0].endswith(",bcra_tea,ai_paragraph")
    assert lineas[1].startswith("2026-10-06,1535.0,1555.0") and lineas[1].endswith("Párrafo de Gemini")
    assert lineas[2].startswith(historico["Fecha"].iloc[0])
    assert len(lineas) == len(historico) + 2


def _series_falsas(hoy=HOY):
    return {
        s.clave: [(hoy - timedelta(days=d), 100.0 + d) for d in (40, 20, 3)]
        for s in agregados.SERIES
    }


def test_series_se_guardan_cuando_existe_la_tabla(entorno):
    deps, hechos, salida = entorno
    guardadas = []
    deps = replace(
        deps,
        tabla_series=lambda engine: True,
        descargar_series=lambda desde: _series_falsas(),
        guardar_series=lambda engine, serie, puntos: guardadas.append(serie.clave) or len(puntos),
    )
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)
    assert _estados(r)["series"] == "ok"
    assert sorted(guardadas) == sorted(s.clave for s in agregados.SERIES)
    assert r.estado == "ok"


def test_una_falla_en_las_series_no_pone_la_corrida_en_rojo(entorno):
    deps, hechos, salida = entorno

    def descargar(desde):
        raise ScraperError("BCRA", "leer agregados monetarios", TimeoutError("timeout"))

    deps = replace(deps, tabla_series=lambda engine: True, descargar_series=descargar)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)
    assert _estados(r)["series"] == "advertencia"
    assert r.estado == "advertencia" and r.exitosa
    assert len(hechos["mails"]) == 2


def _secciones():
    from models import SeccionesIA

    return SeccionesIA(
        resumen="Resumen estructurado del día, con la brecha entre el Blue y el MEP y la tendencia de las 25 ruedas.",
        paralelas="Comentario sobre las cotizaciones paralelas.",
        oficiales="Comentario sobre las cotizaciones oficiales.",
        riesgo_pais="Comentario sobre el riesgo país del día.",
        btc="Comentario sobre BTC y sus medias móviles.",
    )


def test_con_la_columna_ai_secciones_sale_la_llamada_estructurada(entorno):
    deps, hechos, salida = entorno
    guardadas = []
    deps = replace(
        deps,
        columna_secciones=lambda engine: True,
        generar_secciones=lambda prompt: (_secciones(), "gemini-x"),
        guardar_secciones=lambda engine, fecha, secciones, modelo: guardadas.append((fecha, modelo)) or 1,
    )
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert _estados(r)["ia"] == "ok"
    assert guardadas == [(HOY, "gemini-x")]
    assert hechos["parrafos"] == []  # no hizo falta el párrafo único
    html = hechos["mails"][0][0].get_payload()[0].get_payload(decode=True).decode()
    assert "Resumen estructurado del día" in html
    assert "Comentario sobre las cotizaciones paralelas." in html and "Comentario sobre BTC" in html


def test_si_falla_la_estructurada_vuelve_al_parrafo_unico(entorno):
    deps, hechos, salida = entorno
    deps = replace(deps, columna_secciones=lambda engine: True, generar_secciones=lambda prompt: (None, None))
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert _estados(r)["ia"] == "advertencia"
    assert hechos["parrafos"] == [HOY]
    html = hechos["mails"][0][0].get_payload()[0].get_payload(decode=True).decode()
    assert "Párrafo de Gemini" in html and "comentario-ia" not in html


def test_dry_run_con_ia_prueba_los_comentarios_sin_guardar(entorno):
    deps, hechos, salida = entorno
    llamadas, guardadas = [], []
    deps = replace(
        deps,
        generar_secciones=lambda prompt: llamadas.append(prompt) or (_secciones(), "gemini-x"),
        guardar_secciones=lambda *a: guardadas.append(a) or 1,
    )
    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida, probar_ia=True), deps)

    assert _estados(r)["ia"] == "ok" and len(llamadas) == 1 and guardadas == []
    vista = (salida / "mail_preview.html").read_text(encoding="utf-8")
    assert "Comentario sobre el riesgo país del día." in vista


def test_dry_run_sin_con_ia_no_llama_a_gemini(entorno):
    deps, hechos, salida = entorno
    llamadas = []
    deps = replace(deps, generar_secciones=lambda prompt: llamadas.append(prompt) or (None, None))
    pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), deps)
    assert llamadas == []


def test_no_corre_un_feriado(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, feriado=lambda fecha: "Día de prueba"))
    assert r.estado == "omitida"
    assert "feriado (Día de prueba)" in r.etapas[0].detalle
    assert hechos["filas"] == [] and hechos["mails"] == []


def test_no_corre_un_fin_de_semana(entorno):
    deps, hechos, salida = entorno
    sabado = HOY + timedelta(days=4)
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, hoy=lambda: sabado))
    assert r.estado == "omitida" and "fin de semana" in r.etapas[0].detalle


def test_forzar_corre_aunque_sea_feriado(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida, forzar=True), replace(deps, feriado=lambda fecha: "Día de prueba"))
    assert r.estado == "ok" and len(hechos["mails"]) == 2


def test_si_el_calendario_no_responde_se_corre_igual(entorno):
    deps, hechos, salida = entorno

    def feriado(fecha):
        raise ConnectionError("ArgentinaDatos no responde")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, feriado=feriado))
    assert r.estado == "ok" and len(hechos["mails"]) == 2


def test_dry_run_un_feriado_corre_igual(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), replace(deps, feriado=lambda fecha: "Día de prueba"))
    assert r.estado == "ok"


def test_el_log_no_lleva_secretos_ni_destinatarios(tmp_path):
    import logging
    import smtplib

    archivo = tmp_path / "corrida.log"
    pipeline.configurar_logging(archivo)
    try:
        log = logging.getLogger("pipeline.test")
        log.error("FRED respondió 400 para https://api.stlouisfed.org/fred?api_key=fred-falsa&x=1")
        try:
            raise smtplib.SMTPRecipientsRefused({"uno@example.com": (550, b"no existe")})
        except smtplib.SMTPRecipientsRefused:
            log.exception("falló el envío con clave no-es-una-clave")
        for handler in logging.getLogger().handlers:
            handler.flush()
        texto = archivo.read_text(encoding="utf-8")
    finally:
        logging.basicConfig(force=True)

    for sensible in ("fred-falsa", "no-es-una-clave", "uno@example.com", "clave@127.0.0.1"):
        assert sensible not in texto
    assert "api_key=***" in texto and "[destinatario]" in texto


def test_redactar_tapa_la_url_de_la_base_entera():
    url = "postgresql://usuario:clave@127.0.0.1:1/inexistente"
    assert pipeline.redactar(f"no conecta a {url}") == "no conecta a ***"


def test_una_fila_sin_parrafo_no_cuenta_como_dia_hecho(entorno, historico):
    # La corrida anterior insertó la fila y murió antes de la IA y del mail
    deps, hechos, salida = entorno
    con_hoy = historico.copy()
    con_hoy.loc[0, "Fecha"] = str(HOY)
    con_hoy.loc[0, "ai_paragraph"] = None

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, leer_historico=lambda e: (con_hoy, "supabase")))

    assert r.estado == "error" and not r.exitosa
    control = next(e for e in r.etapas if e.nombre == "control")
    assert control.estado == "error" and "no terminó" in control.detalle
    assert hechos["mails"] == [] and len(hechos["alertas"]) == 1


def test_si_el_insert_encuentra_la_fila_no_se_manda_otro_mail(entorno):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, guardar_fila=lambda e, f, sobrescribir=False: False))
    assert r.estado == "omitida" and r.exitosa
    assert hechos["mails"] == [] and hechos["parrafos"] == []


def test_una_excepcion_en_la_ia_estructurada_cae_al_parrafo_unico(entorno):
    deps, hechos, salida = entorno

    def generar(prompt):
        raise TimeoutError("Gemini no respondió")

    deps = replace(deps, columna_secciones=lambda engine: True, generar_secciones=generar)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert _estados(r)["ia"] == "advertencia"
    assert hechos["parrafos"] == [HOY] and hechos["limpiezas"] == [HOY]
    html = hechos["mails"][0][0].get_payload()[0].get_payload(decode=True).decode()
    assert "Párrafo de Gemini" in html


def test_si_no_se_pueden_guardar_las_secciones_cae_al_parrafo_unico(entorno):
    deps, hechos, salida = entorno

    def guardar(engine, fecha, secciones, modelo):
        raise ConnectionError("Supabase no responde")

    deps = replace(deps, columna_secciones=lambda engine: True,
                   generar_secciones=lambda prompt: (_secciones(), "gemini-x"), guardar_secciones=guardar)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)
    assert _estados(r)["ia"] == "advertencia" and hechos["parrafos"] == [HOY]


def test_una_serie_rota_no_frena_a_las_demas(entorno):
    deps, hechos, salida = entorno
    series = _series_falsas()
    series["m2"] = []  # discontinuada
    guardadas = []
    deps = replace(deps, tabla_series=lambda engine: True, descargar_series=lambda desde: series,
                   guardar_series=lambda engine, serie, puntos: guardadas.append(serie.clave) or len(puntos))

    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert _estados(r)["series"] == "advertencia"
    assert "m2" not in guardadas and len(guardadas) == len(agregados.SERIES) - 1


@pytest.mark.parametrize("argumentos", [
    ["--enviar-a", "yo@example.com"],
    ["--con-ia"],
    ["--dry-run", "--sin-mail", "--enviar-a", "yo@example.com"],
])
def test_combinaciones_de_opciones_que_no_se_admiten(argumentos):
    with pytest.raises(SystemExit):
        pipeline.main(argumentos)


# ── Etapa previews ───────────────────────────────────────────────────────────

def test_una_corrida_real_actualiza_previews(entorno):
    deps, hechos, salida = entorno
    llamadas = []
    deps = replace(deps, actualizar_previews=lambda repo, archivos=None: llamadas.append((repo, archivos)) or (True, "ok"))

    r = pipeline.correr(pipeline.Opciones(), deps)

    assert _estados(r)["previews"] == "ok"
    # A main van solo los gráficos; el PowerPoint del día va a su propia rama
    assert llamadas == [(pipeline.RAIZ, list(charts.ORDEN_EN_MAIL))]
    assert {p.name for p in pipeline.PREVIEWS.glob("*.jpg")} == set(charts.ORDEN_EN_MAIL)
    assert hechos["presentaciones"] == [(pipeline.PREVIEWS / presentacion.ARCHIVO, "Reporte ejecutivo del 06/10/2026")]


def test_previews_sin_cambios_queda_omitida(entorno):
    deps, _, _ = entorno
    deps = replace(deps, actualizar_previews=lambda repo, archivos=None: (False, "sin cambios"),
                   publicar_presentacion=lambda repo, archivo, mensaje=None: (False, "git no está instalado: se omite"))
    r = pipeline.correr(pipeline.Opciones(), deps)
    assert _estados(r)["previews"] == "omitida" and r.estado == "ok"


def test_un_push_que_falla_pone_la_corrida_en_rojo_con_una_alerta(entorno):
    import preview_git

    deps, hechos, _ = entorno

    def actualizar(repo, archivos=None):
        raise preview_git.GitError("git push terminó con código 128")

    r = pipeline.correr(pipeline.Opciones(), replace(deps, actualizar_previews=actualizar))
    assert _estados(r)["previews"] == "error" and r.estado == "error"
    assert len(hechos["mails"]) == 2 and len(hechos["alertas"]) == 1


def test_previews_no_se_toca_si_ruta_repo_es_otra_carpeta(entorno, monkeypatch, tmp_path):
    deps, hechos, _ = entorno
    monkeypatch.setattr(pipeline.settings, "ruta_repo", tmp_path / "otro-repo")
    r = pipeline.correr(pipeline.Opciones(), deps)
    previews = next(e for e in r.etapas if e.nombre == "previews")
    assert previews.estado == "omitida" and "RUTA_REPO" in previews.detalle


# ── Gráficos, dry-run y otras ramas ──────────────────────────────────────────

def test_un_grafico_que_falla_no_viaja_aunque_quede_el_de_ayer(entorno, monkeypatch):
    deps, hechos, salida = entorno
    (salida / charts.INFLACION).write_bytes(b"jpg de ayer")

    def falla(*a, **k):
        raise ValueError("datos raros")

    monkeypatch.setattr(pipeline.charts, "grafico_inflacion", falla)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    assert _estados(r)["graficos"] == "error" and r.estado == "error"
    mensaje = hechos["mails"][0][0]
    cids = [p.get("Content-ID") for p in mensaje.walk() if p.get_content_type() == "image/jpeg"]
    assert cids == ["<image1>", "<image3>", "<image4>"]
    assert len(hechos["alertas"]) == 1


def test_dry_run_con_la_fila_de_hoy_usa_el_parrafo_y_los_comentarios_guardados(entorno, historico):
    deps, hechos, salida = entorno
    con_hoy = historico.copy()
    con_hoy.loc[0, "Fecha"] = str(HOY)
    con_hoy.loc[0, "ai_paragraph"] = "Párrafo guardado del día"
    con_hoy["ai_secciones"] = None
    con_hoy.at[0, "ai_secciones"] = {"paralelas": "Comentario guardado de paralelas.", "btc": None}

    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), replace(deps, leer_historico=lambda e: (con_hoy, "supabase")))

    assert r.estado == "ok"
    vista = (salida / "mail_preview.html").read_text(encoding="utf-8")
    assert "Párrafo guardado del día" in vista and "Comentario guardado de paralelas." in vista


def test_si_no_se_puede_crear_el_engine_alerta(entorno):
    deps, hechos, salida = entorno

    def crear():
        raise ValueError("URL de base inválida")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, crear_engine=crear))
    assert _estados(r) == {"pipeline": "error"} and len(hechos["alertas"]) == 1


def test_si_el_candado_falla_la_corrida_sigue(entorno):
    deps, hechos, salida = entorno

    def candado(engine):
        raise ConnectionError("el pooler no responde")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, candado=candado))
    assert r.estado == "ok" and len(hechos["mails"]) == 2


def test_sin_mail_y_sin_push(entorno):
    deps, hechos, _ = entorno
    r = pipeline.correr(pipeline.Opciones(enviar_mail=False, push_previews=False), deps)
    assert _estados(r)["mail"] == "omitida" and _estados(r)["previews"] == "omitida"
    assert hechos["mails"] == [] and hechos["previews"] == [] and len(hechos["filas"]) == 1


def test_enviar_a_manda_un_solo_mail_con_el_csv(entorno):
    deps, hechos, salida = entorno
    pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida, enviar_a=["yo@example.com"]), deps)
    mensaje, _ = hechos["mails"][0]
    assert any(p.get_content_type() == "text/csv" for p in mensaje.walk())


def test_secciones_guardadas_en_cero_filas_caen_al_parrafo_unico(entorno):
    deps, hechos, salida = entorno
    deps = replace(deps, columna_secciones=lambda engine: True,
                   generar_secciones=lambda prompt: (_secciones(), "gemini-x"),
                   guardar_secciones=lambda engine, fecha, secciones, modelo: 0)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)
    assert _estados(r)["ia"] == "advertencia" and hechos["parrafos"] == [HOY]


def test_si_no_se_puede_consultar_la_columna_sale_el_parrafo_unico(entorno):
    deps, hechos, salida = entorno

    def columna(engine):
        raise ConnectionError("Supabase no responde")

    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, columna_secciones=columna))
    assert _estados(r)["ia"] == "ok" and hechos["parrafos"] == [HOY]


def test_el_json_de_la_corrida_tiene_lo_que_lee_app(entorno, tmp_path, monkeypatch):
    """El contrato entre pipeline.py y app.py: --json con estado y etapas[nombre, estado, segundos]."""

    from fastapi.testclient import TestClient

    import app

    deps, _, salida = entorno
    resultado = pipeline.correr(pipeline.Opciones(salida=salida), deps)
    monkeypatch.setattr(pipeline, "correr", lambda opciones: resultado)
    monkeypatch.setattr(pipeline, "configurar_logging", lambda archivo: None)
    escrito = tmp_path / "resultado.json"
    pipeline.main(["--dry-run", "--json", str(escrito)])

    def run(comando, **kwargs):
        destino = comando[comando.index("--json") + 1]
        Path(destino).write_text(escrito.read_text(encoding="utf-8"), encoding="utf-8")
        return subprocess.CompletedProcess(comando, 0)

    monkeypatch.setattr(app.subprocess, "run", run)
    datos = TestClient(app.app).post("/run", headers={"x-api-key": "clave-de-test"}).json()

    assert datos["estado"] == "ok"
    assert [e["nombre"] for e in datos["etapas"]] == [e.nombre for e in resultado.etapas]
    assert all(set(e) == {"nombre", "estado", "segundos"} for e in datos["etapas"])


def test_historico_desde_el_csv_es_una_advertencia(entorno, historico):
    deps, hechos, salida = entorno
    r = pipeline.correr(pipeline.Opciones(salida=salida), replace(deps, leer_historico=lambda e: (historico.copy(), "csv")))
    assert _estados(r)["historico"] == "advertencia" and r.estado == "advertencia"
    assert len(hechos["mails"]) == 2


def test_la_guarda_de_conftest_corta_las_fuentes_reales():
    """Si alguien saca la guarda, este test lo nota: el default de Dependencias saldría a Yahoo."""
    from scrapers.utils import ScraperError

    with pytest.raises(ScraperError, match="un test intentó bajar datos de Yahoo") as error:
        pipeline.Dependencias().descargar_btc(HOY - timedelta(days=30), HOY)
    assert isinstance(error.value.__cause__, AssertionError)


# ── Etapa presentacion ───────────────────────────────────────────────────────

def _diapositivas(ruta):
    from pptx import Presentation

    return [" ".join(s.text_frame.text for s in d.shapes if s.has_text_frame) for d in Presentation(ruta).slides]


def test_dry_run_arma_la_presentacion_en_la_salida_y_no_en_previews(entorno):
    deps, _, salida = entorno
    r = pipeline.correr(pipeline.Opciones(dry_run=True, salida=salida), deps)

    assert _estados(r)["presentacion"] == "ok"
    assert len(_diapositivas(salida / presentacion.ARCHIVO)) == 6
    assert not (pipeline.PREVIEWS / presentacion.ARCHIVO).exists()


def test_sin_push_igual_arma_la_presentacion(entorno):
    deps, hechos, _ = entorno
    r = pipeline.correr(pipeline.Opciones(push_previews=False), deps)
    assert _estados(r)["presentacion"] == "ok" and _estados(r)["previews"] == "omitida"
    assert (pipeline.PREVIEWS / presentacion.ARCHIVO).exists() and hechos["previews"] == []


def test_la_presentacion_lleva_los_textos_de_esta_corrida(entorno):
    """El párrafo y los comentarios salen de la IA de hoy, en memoria: la fila no los tiene."""
    deps, _, salida = entorno
    deps = replace(deps, columna_secciones=lambda engine: True,
                   generar_secciones=lambda prompt: (_secciones(), "gemini-x"),
                   guardar_secciones=lambda engine, fecha, secciones, modelo: 1)

    pipeline.correr(pipeline.Opciones(salida=salida), deps)

    textos = _diapositivas(salida / presentacion.ARCHIVO)
    assert _secciones().resumen in textos[2]
    assert _secciones().paralelas in textos[3] and _secciones().btc in textos[5]


def test_si_la_presentacion_falla_es_advertencia_y_el_mail_sale(entorno, monkeypatch):
    deps, hechos, _ = entorno
    llamadas = []
    deps = replace(deps, actualizar_previews=lambda repo, archivos=None: llamadas.append(archivos) or (True, "ok"))

    def falla(*args, **kwargs):
        raise ValueError("plantilla rota")

    monkeypatch.setattr(pipeline.presentacion, "armar", falla)
    r = pipeline.correr(pipeline.Opciones(), deps)

    etapa = next(e for e in r.etapas if e.nombre == "presentacion")
    assert etapa.estado == "advertencia" and "plantilla rota" in etapa.detalle
    assert r.estado == "advertencia" and r.exitosa
    assert len(hechos["mails"]) == 2 and hechos["alertas"] == []
    # Los gráficos se publican igual, y no se sube ningún PowerPoint
    assert llamadas == [list(charts.ORDEN_EN_MAIL)] and hechos["presentaciones"] == []


def test_la_presentacion_avisa_el_grafico_que_falto(entorno, monkeypatch):
    deps, _, salida = entorno

    def falla(*a, **k):
        raise ValueError("datos raros")

    monkeypatch.setattr(pipeline.charts, "grafico_inflacion", falla)
    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    etapa = next(e for e in r.etapas if e.nombre == "presentacion")
    assert etapa.estado == "ok" and charts.INFLACION in etapa.detalle
    assert (salida / presentacion.ARCHIVO).exists()


def test_si_hoy_no_se_armo_la_presentacion_no_se_sube_la_de_ayer(entorno, monkeypatch):
    deps, hechos, _ = entorno
    pipeline.PREVIEWS.mkdir(parents=True, exist_ok=True)
    (pipeline.PREVIEWS / presentacion.ARCHIVO).write_bytes(b"la de ayer")

    def falla(*args, **kwargs):
        raise ValueError("plantilla rota")

    monkeypatch.setattr(pipeline.presentacion, "armar", falla)
    pipeline.correr(pipeline.Opciones(), deps)
    assert hechos["presentaciones"] == []


def test_si_falla_la_publicacion_de_la_presentacion_es_advertencia(entorno):
    deps, hechos, _ = entorno

    def publicar(repo, archivo, mensaje=None):
        raise RuntimeError("push rechazado")

    r = pipeline.correr(pipeline.Opciones(), replace(deps, publicar_presentacion=publicar))

    etapa = next(e for e in r.etapas if e.nombre == "previews")
    assert etapa.estado == "advertencia" and "push rechazado" in etapa.detalle
    assert r.estado == "advertencia" and r.exitosa
    assert len(hechos["previews"]) == 1 and len(hechos["mails"]) == 2 and hechos["alertas"] == []


def test_sin_cambios_en_graficos_pero_con_presentacion_publicada_queda_ok(entorno):
    deps, hechos, _ = entorno
    r = pipeline.correr(pipeline.Opciones(), replace(deps, actualizar_previews=lambda repo, archivos=None: (False, "sin cambios")))
    etapa = next(e for e in r.etapas if e.nombre == "previews")
    assert etapa.estado == "ok" and "publicada" in etapa.detalle


def test_sin_python_pptx_la_corrida_sigue_y_avisa(entorno, monkeypatch):
    deps, hechos, salida = entorno
    monkeypatch.setattr(pipeline, "presentacion", None)
    monkeypatch.setattr(pipeline, "FALTA_PRESENTACION", "ModuleNotFoundError: No module named 'pptx'", raising=False)

    r = pipeline.correr(pipeline.Opciones(salida=salida), deps)

    etapa = next(e for e in r.etapas if e.nombre == "presentacion")
    assert etapa.estado == "advertencia" and "requirements.txt" in etapa.detalle
    assert r.exitosa and len(hechos["mails"]) == 2


def test_dry_run_con_salida_dentro_de_previews_usa_un_temporal(monkeypatch, tmp_path):
    monkeypatch.setattr(pipeline, "PREVIEWS", tmp_path / "Previews")
    carpeta = pipeline._carpeta_de_salida(pipeline.Opciones(dry_run=True, salida=tmp_path / "Previews" / "prueba"))
    assert not carpeta.resolve().is_relative_to(tmp_path / "Previews")
